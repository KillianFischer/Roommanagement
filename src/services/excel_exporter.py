from typing import List, Dict, Tuple, Optional
import os
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from tkinter import messagebox

from models.student import StudentPreference


class ExcelExporter:
    def __init__(self):

        self.header_fill = PatternFill(start_color="2F5596", end_color="2F5596", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF")
        self.subheader_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
        self.subheader_font = Font(bold=True)
        self.alt_row_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def export_student_schedules(self, filepath: str, schedule: Dict[Tuple[str, int], any], 
                                student_preferences: List[StudentPreference], time_slots: List[Tuple[str, str]]):
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Schüler Zeitpläne"
        
        ws.cell(row=1, column=1, value="Schüler Zeitpläne")
        ws.cell(row=1, column=1).font = Font(bold=True, size=16)
        ws.cell(row=2, column=1, value=f"Erstellt am: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        
        student_schedules = self._prepare_student_schedules(schedule, student_preferences)
        
        current_row = 4
        
        sorted_students = sorted(student_schedules.keys(), key=lambda x: x.lower())
        
        for student_name in sorted_students:
            ws.cell(row=current_row, column=1, value=f"Schüler: {student_name}")
            ws.cell(row=current_row, column=1).font = self.subheader_font
            current_row += 1
            
            headers = ["Slot", "Zeit", "Unternehmen", "Raum"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.alignment = Alignment(horizontal='center')
                cell.border = self.thin_border
            
            current_row += 1
            
            student_data = student_schedules[student_name]
            
            for row_idx, (slot_idx, (slot_letter, time_range)) in enumerate(enumerate(time_slots), 0):
                row_fill = self.alt_row_fill if row_idx % 2 == 0 else None
                
                cell = ws.cell(row=current_row, column=1, value=slot_letter)
                if row_fill:
                    cell.fill = row_fill
                cell.border = self.thin_border
                
                cell = ws.cell(row=current_row, column=2, value=time_range)
                if row_fill:
                    cell.fill = row_fill
                cell.border = self.thin_border
                
                if slot_idx in student_data:
                    company_name, room = student_data[slot_idx]
                    
                    cell = ws.cell(row=current_row, column=3, value=company_name)
                    if row_fill:
                        cell.fill = row_fill
                    cell.border = self.thin_border
                    
                    cell = ws.cell(row=current_row, column=4, value=room)
                    if row_fill:
                        cell.fill = row_fill
                    cell.border = self.thin_border
                else:
                    cell = ws.cell(row=current_row, column=3, value="-")
                    if row_fill:
                        cell.fill = row_fill
                    cell.border = self.thin_border
                    
                    cell = ws.cell(row=current_row, column=4, value="-")
                    if row_fill:
                        cell.fill = row_fill
                    cell.border = self.thin_border
                
                current_row += 1
            
            current_row += 1
        
        ws.column_dimensions['A'].width = 10  # Slot
        ws.column_dimensions['B'].width = 20  # Zeit
        ws.column_dimensions['C'].width = 40  # Unternehmen
        ws.column_dimensions['D'].width = 20  # Raum
        
        wb.save(filepath)
        return True
            
    def _prepare_student_schedules(self, schedule, student_preferences):
        student_schedules = {}
        
        for student in student_preferences:
            student_schedules[student.name] = {}
            
        for (company_id, slot_idx), session in schedule.items():
            if slot_idx == -1:
                continue  # Skip excluded companies
                
            for student in session.students:
                student_name = student["name"]
                if student_name not in student_schedules:
                    student_schedules[student_name] = {}
                    
                student_schedules[student_name][slot_idx] = (session.company.name, session.room)
                
        return student_schedules
        
    def export_company_overview(self, filepath: str, schedule: Dict[Tuple[str, int], any],
                               time_slots: List[Tuple[str, str]]):
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Unternehmensübersicht"
        
        ws.cell(row=1, column=1, value="Unternehmensübersicht")
        ws.cell(row=1, column=1).font = Font(bold=True, size=16)
        ws.cell(row=2, column=1, value=f"Erstellt am: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        
        slot_to_companies = {}
        for (company_id, slot_idx), session in schedule.items():
            if slot_idx == -1:
                continue  # Skip excluded companies
                
            if slot_idx not in slot_to_companies:
                slot_to_companies[slot_idx] = []
                
            slot_to_companies[slot_idx].append((session.company.name, session.room, len(session.students)))
        
        current_row = 4
        
        for slot_idx, (slot_letter, time_range) in enumerate(time_slots):
            if slot_idx not in slot_to_companies:
                continue
                
            ws.cell(row=current_row, column=1, value=f"Slot {slot_letter}: {time_range}")
            ws.cell(row=current_row, column=1).font = self.subheader_font
            current_row += 1
            
            headers = ["Unternehmen", "Raum", "Anzahl Schüler"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.alignment = Alignment(horizontal='center')
                cell.border = self.thin_border
            
            current_row += 1
            
            for row_idx, (company_name, room, student_count) in enumerate(
                sorted(slot_to_companies[slot_idx], key=lambda x: x[0].lower())
            ):
                row_fill = self.alt_row_fill if row_idx % 2 == 0 else None
                
                cell = ws.cell(row=current_row, column=1, value=company_name)
                if row_fill:
                    cell.fill = row_fill
                cell.border = self.thin_border
                
                cell = ws.cell(row=current_row, column=2, value=room)
                if row_fill:
                    cell.fill = row_fill
                cell.border = self.thin_border
                
                cell = ws.cell(row=current_row, column=3, value=student_count)
                if row_fill:
                    cell.fill = row_fill
                cell.alignment = Alignment(horizontal='center')
                cell.border = self.thin_border
                
                current_row += 1
            
            current_row += 1
        
        excluded_companies = []
        for (company_id, slot_idx), session in schedule.items():
            if slot_idx == -1:
                excluded_companies.append(session.company.name)
                
        if excluded_companies:
            ws.cell(row=current_row, column=1, value="Ausgeschlossene Unternehmen")
            ws.cell(row=current_row, column=1).font = self.subheader_font
            current_row += 1
            
            for company_name in sorted(excluded_companies):
                ws.cell(row=current_row, column=1, value=f"• {company_name}: Kein Schülerinteresse")
                current_row += 1
        
        ws.column_dimensions['A'].width = 40  # Unternehmen
        ws.column_dimensions['B'].width = 20  # Raum
        ws.column_dimensions['C'].width = 15  # Anzahl Schüler
        
        wb.save(filepath)
        return True

    def export_attendance_lists(self, filepath: str, schedule: Dict[Tuple[str, int], any], 
                              time_slots: List[Tuple[str, str]], preview_mode=False):
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Anwesenheitslisten"
        
        sorted_sessions = sorted(
            schedule.items(),
            key=lambda x: (x[0][0], x[0][1]),
        )
        
        if preview_mode:
            company_ids = list(
                set(company_id for (company_id, _), _ in sorted_sessions)
            )
            if len(company_ids) > 6:
                company_ids = company_ids[:6]
            sorted_sessions = [
                (key, session)
                for (key, session) in sorted_sessions
                if key[0] in company_ids
            ]
        
        current_row = 1
        
        for (company_id, slot_idx), session in sorted_sessions:
            if slot_idx == -1:
                continue
                
            slot_letter, time_range = time_slots[slot_idx]
            
            ws.cell(row=current_row, column=1, value=session.company.name)
            ws.cell(row=current_row, column=1).font = Font(bold=True, size=14)
            current_row += 1
            
            ws.cell(row=current_row, column=1, value=f"Zeitfenster: {slot_letter} ({time_range})")
            current_row += 1
            
            ws.cell(row=current_row, column=1, value=f"Raum: {session.room}")
            current_row += 1
            
            headers = ["Nr.", "Name", "Klasse", "Anwesend"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.alignment = Alignment(horizontal='center')
                cell.border = self.thin_border
            
            current_row += 1
            
            if len(session.students) == 0:
                cell = ws.cell(row=current_row, column=2, value="Kein Schülerinteresse")
                cell.border = self.thin_border
                current_row += 1
            else:
                for i, student in enumerate(sorted(session.students, key=lambda x: x["name"]), 1):
                    class_name = student["id"].split("_")[0]
                    
                    cell = ws.cell(row=current_row, column=1, value=i)
                    cell.border = self.thin_border
                    cell.alignment = Alignment(horizontal='center')
                    
                    cell = ws.cell(row=current_row, column=2, value=student["name"])
                    cell.border = self.thin_border
                    
                    cell = ws.cell(row=current_row, column=3, value=class_name)
                    cell.border = self.thin_border
                    cell.alignment = Alignment(horizontal='center')
                    
                    cell = ws.cell(row=current_row, column=4, value="")
                    cell.border = self.thin_border
                    
                    current_row += 1
            
            current_row += 2
        
        ws.column_dimensions['A'].width = 5   # Nr.
        ws.column_dimensions['B'].width = 30  # Name
        ws.column_dimensions['C'].width = 15  # Klasse
        ws.column_dimensions['D'].width = 15  # Anwesend
        
        wb.save(filepath)
        return True
        
    def export_room_list(self, filepath: str, rooms: List[str], room_capacities: Dict[str, int]):
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Raumliste"
        
        ws.cell(row=1, column=1, value="Raumliste")
        ws.cell(row=1, column=1).font = Font(bold=True, size=16)
        ws.cell(row=2, column=1, value=f"Erstellt am: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        
        current_row = 4
        
        headers = ["Raum", "Kapazität"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.thin_border
        
        current_row += 1
        
        for row_idx, room_name in enumerate(sorted(rooms)):
            row_fill = self.alt_row_fill if row_idx % 2 == 0 else None
            
            cell = ws.cell(row=current_row, column=1, value=room_name)
            if row_fill:
                cell.fill = row_fill
            cell.border = self.thin_border
            
            capacity = room_capacities.get(room_name, 30)  # Default to 30 if not specified
            cell = ws.cell(row=current_row, column=2, value=capacity)
            if row_fill:
                cell.fill = row_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.thin_border
            
            current_row += 1
        
        ws.column_dimensions['A'].width = 30  # Raum
        ws.column_dimensions['B'].width = 15  # Kapazität
        
        wb.save(filepath)
        return True

    def export_schedule(self, filepath: str, schedule: dict, time_slots: list, companies: list) -> bool:
        """Exports the main schedule grid view to an Excel file."""
        try:
            header = ["Unternehmen"] + [f"{slot} ({time})" for slot, time in time_slots]
            data = []

            for company in companies:
                # Skip companies that have been excluded (assuming an entry with slot_idx -1 exists)
                if any((company.unique_id, -1) == key for key in schedule.keys()):
                    continue
                    
                display_name = str(company)  # Use the __str__ representation
                row = [display_name]

                for slot_idx, _ in enumerate(time_slots):
                    if slot_idx < company.earliest_slot or slot_idx in company.blocked_slots:
                        text = "---" # Indicate blocked or too early slot
                    else:
                        session = schedule.get((company.unique_id, slot_idx))
                        if session:
                            text = f"Raum {session.room}" # Show room
                        else:
                            text = "---" # Indicate no session scheduled
                    row.append(text)
                data.append(row)

            df = pd.DataFrame(data, columns=header)
            df.to_excel(filepath, index=False)
            return True
            
        except Exception as e:
            messagebox.showerror("Excel Export Fehler", f"Fehler beim Exportieren des Zeitplans nach Excel: {str(e)}")
            return False 